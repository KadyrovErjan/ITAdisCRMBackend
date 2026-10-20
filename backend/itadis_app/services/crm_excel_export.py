"""Operational CRM workbook export.

This exporter deliberately reads only confirmed CRM entities.  Staging rows and
``HistoricalReceipt`` values are excluded because they are not current cash.
"""
from __future__ import annotations

from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from django.db.models import Prefetch
from django.utils import timezone

from ..models import Transaction
from .payment_plans import financial_summary


MONEY_FORMAT = '#,##0.00'
HEADER_FILL = PatternFill('solid', fgColor='D9EAD3')


def _header(sheet, values):
    sheet.append(values)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
    sheet.freeze_panes = 'A2'


def build_crm_groups_export(groups):
    """Return an XLSX workbook for the groups visible to the requesting user."""
    workbook = Workbook()
    groups_sheet = workbook.active
    groups_sheet.title = 'Топтор'
    _header(groups_sheet, [
        'Топ', 'Технология', 'Расписание', 'Дата начала', 'Длительность, ай',
        'Количество учеников',
    ])

    students_sheet = workbook.create_sheet('Окуучулар')
    _header(students_sheet, [
        'Топ', 'Окуучу', 'Телефон', 'Курстун баасы', 'Төлөнгөн',
        'Келишим боюнча калдык', 'Кечиккен карыз', 'Кийинки төлөм күнү',
        'Кийинки төлөм суммасы', 'Окуу абалы', 'Төлөм абалы',
    ])

    payments_sheet = workbook.create_sheet('Тастыкталган төлөмдөр')
    _header(payments_sheet, ['Топ', 'Окуучу', 'Күнү', 'Түрү', 'Сумма', 'Кассир'])

    for group in groups:
        students = list(group.students.select_related('registered_by').prefetch_related(
            Prefetch('transactions', queryset=Transaction.objects.select_related('cashier').order_by('created_at')),
        ).order_by('full_name'))
        groups_sheet.append([
            group.name, group.technology or group.subject, group.schedule,
            group.start_date, group.duration_months, len(students),
        ])
        for student in students:
            summary = financial_summary(student)
            students_sheet.append([
                group.name, student.full_name, student.phone or '', student.course_price,
                Decimal(summary['total_paid']), summary['contract_remaining'],
                summary['overdue_amount'], summary['next_payment_date'],
                summary['next_payment_amount'], student.effective_learning_status,
                summary['payment_status'],
            ])
            for receipt in student.transactions.all():
                payments_sheet.append([
                    group.name, student.full_name, timezone.localtime(receipt.created_at).date(),
                    receipt.type, receipt.amount, receipt.cashier.full_name,
                ])

    for sheet in workbook.worksheets:
        for column in sheet.columns:
            width = min(max((len(str(cell.value or '')) for cell in column), default=10) + 2, 40)
            sheet.column_dimensions[column[0].column_letter].width = width
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, Decimal):
                    cell.number_format = MONEY_FORMAT
    return workbook
